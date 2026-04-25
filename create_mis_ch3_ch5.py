from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\gerla\Downloads\MIS-FINAL-1-3.docx")
OUTPUT = Path(r"C:\Users\gerla\Downloads\MIS-FINAL-1-3-Chapters-3-to-5.docx")


def set_run_font(run, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=False):
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5) if first_line else None
    for run in paragraph.runs:
        set_run_font(run, bool(run.bold))


def add_para(doc, text="", *, bold=False, center=False, first_line=True):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    format_paragraph(
        paragraph,
        align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line=first_line and not center,
    )
    return paragraph


def add_section_heading(doc, text):
    return add_para(doc, text, bold=True, center=False, first_line=False)


def add_chapter(doc, number, title):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_para(doc, number, bold=True, center=True, first_line=False)
    add_para(doc, title, bold=True, center=True, first_line=False)
    add_para(doc, "", first_line=False)


def set_cell(cell, text, bold=False, center=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def add_schedule_table(doc):
    add_para(doc, "Development and Implementation Schedule", center=True, first_line=False)
    add_para(doc, "Planning and Analysis Phase (Week 1-2)", bold=True, first_line=False)
    add_para(doc, "During this phase, the proponents finalize system requirements, user roles, data fields, workflow rules, and reporting needs for research and extension records. This phase establishes the basis for the system design and ensures that the proposed MIS reflects the actual needs of the Research and Extension Office.")
    add_para(doc, "System Design Phase (Week 3-4)", bold=True, first_line=False)
    add_para(doc, "The system design phase covers the preparation of the database structure, interface layout, process flow, security rules, and module design. The design includes the research module, extension module, budget and expense reporting, document upload, dashboard summaries, and automated reports.")
    add_para(doc, "Development Phase (Week 5-9)", bold=True, first_line=False)
    add_para(doc, "The development phase focuses on creating the login and user management module, research management module, extension management module, budget and expense reporting module, document upload function, dashboard, and report generation tools. Each module is developed according to the approved system design.")
    add_para(doc, "Testing and Revision Phase (Week 10-11)", bold=True, first_line=False)
    add_para(doc, "During this phase, the proponents conduct unit testing, integration testing, system testing, usability checking, and user acceptance testing with selected faculty members and office personnel. Errors and usability concerns are corrected before deployment.")
    add_para(doc, "Deployment and Training Phase (Week 12)", bold=True, first_line=False)
    add_para(doc, "The deployment and training phase includes installing the system on the target server, creating official user accounts, migrating available records, orienting users, and monitoring initial operation.")
    add_para(doc, "Post-Implementation Support Phase (Week 13 onward)", bold=True, first_line=False)
    add_para(doc, "After deployment, the proponents collect user feedback, fix minor errors, verify backup procedures, update documentation, and prepare recommended improvements for future versions of the system.")


def normalize_existing_chapter3(doc):
    replacements = {
        "CHAPTER 3": "CHAPTER III",
        "SNLC in Waterfall Methodology": "SDLC in Waterfall Methodology",
        "This chapter researched methodology: Requirements Analysis, System Design, Implementation, Testing, Deployment of System, Maintenance, Research Design, and Subject of the Study.": (
            "This chapter presents the system design and methodology used in the development of the Web-Based Management Information System for Research and Extension of J.H. Cerilles State College - Dumingag Campus. It discusses the SDLC Waterfall Methodology, requirements analysis, system design, implementation, testing, deployment, maintenance, research design, and subjects of the study."
        ),
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
        if paragraph.text.strip() in {"CHAPTER III", "SYSTEM DESIGN AND METHODOLOGY"}:
            for run in paragraph.runs:
                set_run_font(run, bold=True)
            format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)


def main():
    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    normalize_existing_chapter3(doc)

    add_section_heading(doc, "Maintenance")
    add_para(doc, "After deployment, the Management Information System will undergo continuous maintenance to ensure that it remains functional, secure, and useful to the Research and Extension Office. The proponents will monitor system performance, review user feedback, correct minor errors, and update system components when necessary. Regular maintenance will also include checking user accounts, validating uploaded records, reviewing backup files, and ensuring that only authorized users can access sensitive research and extension documents.")
    add_para(doc, "The maintenance phase is important because the needs of the institution may change over time. New report formats, additional research categories, or revised extension documentation requirements may be added in future versions. Through scheduled review and proper technical support, the system can continue to support accurate monitoring, faster retrieval of records, and reliable report generation for accreditation and administrative decision-making.")

    add_section_heading(doc, "Research Design")
    add_para(doc, "This study uses a developmental research design supported by descriptive procedures. The developmental aspect focuses on designing, building, and validating a web-based Management Information System for Research and Extension. The descriptive aspect is used to identify the current problems of the manual process, describe the needs of the users, and determine whether the developed system satisfies the requirements of the intended stakeholders.")
    add_para(doc, "The Waterfall Model serves as the system development guide because it follows a clear sequence of phases: requirements analysis, system design, development, testing, deployment, and maintenance. This model is appropriate for the project because the core requirements of the Research and Extension Office can be identified before implementation. Each phase produces outputs that guide the next phase, helping the proponents maintain organized documentation and systematic development.")

    add_section_heading(doc, "Subject of the Study")
    add_para(doc, "The subjects of the study are the intended users and beneficiaries of the Web-Based Management Information System for Research and Extension of J.H. Cerilles State College - Dumingag Campus. These include the Research and Extension Office personnel, faculty researchers, extension proponents, administrators, and selected staff members who are involved in submitting, validating, monitoring, and reporting research and extension activities.")
    add_para(doc, "Faculty members serve as primary users because they will encode research information, upload abstracts and supporting documents, submit extension activity records, and update accomplishment reports. The Research and Extension Office personnel and administrators serve as validating and monitoring users because they will review submissions, manage records, track project status, and generate reports needed for institutional assessment, accreditation, and planning.")

    add_section_heading(doc, "Data Gathering Procedure")
    add_para(doc, "The proponents gathered data by reviewing the existing documentation practices of the Research and Extension Office and identifying the difficulties experienced in manual recordkeeping. Interviews and informal discussions were conducted with personnel who handle research and extension records to determine the required data fields, approval flow, report formats, and document management needs. Observation of current processes also helped the proponents understand how proposals, accomplishments, budgets, and supporting documents are stored and retrieved.")
    add_para(doc, "The gathered information served as the basis for defining the functional and non-functional requirements of the system. These requirements guided the design of the database, user roles, interface flow, security controls, and report generation features. After development, feedback from selected users will be used to evaluate the usability, usefulness, and completeness of the system.")

    add_chapter(doc, "CHAPTER IV", "DEVELOPMENT, TESTING AND IMPLEMENTATION")
    add_section_heading(doc, "Description of the Prototype")
    add_para(doc, "The developed prototype is a web-based Management Information System for Research and Extension designed for J.H. Cerilles State College - Dumingag Campus. It provides a centralized platform where faculty members, research and extension personnel, and administrators can manage research outputs, extension activities, supporting documents, budgets, and reports. The system is intended to reduce the reliance on scattered paper files, spreadsheets, and manually prepared reports.")
    add_para(doc, "The prototype includes a secure login module that controls access according to user roles. Faculty users can encode research proposals, completed studies, presented papers, published works, utilized outputs, and extension activity records. They can also upload abstracts, certificates, conference proceedings, narrative reports, attendance records, and other supporting files. Research and Extension Office personnel can validate submissions, monitor project status, review accomplishments, and manage records submitted by faculty proponents.")
    add_para(doc, "The system also includes modules for extension target planning, actual accomplishment recording, budget utilization, expense report submission, document storage, and automated report generation. Through the dashboard, authorized users can view summaries of completed, ongoing, presented, published, and utilized research outputs, as well as planned and accomplished extension activities. These features support faster retrieval of information and more organized documentation for accreditation, institutional assessment, and administrative decision-making.")

    add_section_heading(doc, "Gantt Chart")
    add_schedule_table(doc)
    add_para(doc, "The schedule presents the planned development and implementation activities for the Web-Based Management Information System for Research and Extension. It begins with requirements gathering and system design, followed by module development, testing, deployment, training, and post-implementation support. The timeline provides a guide for organizing project activities and ensuring that the system is developed, tested, and introduced to users in a systematic manner.")

    add_section_heading(doc, "Development")
    add_para(doc, "The system will be developed as a modular web application using commonly available web technologies. PHP will be used for server-side processing, while MySQL or MariaDB will be used to store research records, extension records, user accounts, uploaded document details, budget entries, and generated reports. The user interface will be created using HTML, CSS, and JavaScript to provide a clean and responsive layout that can be accessed through standard web browsers.")
    add_para(doc, "The development process will follow the system design prepared in Chapter III. Separate modules will be created for authentication, user management, research management, extension management, budget and expense reporting, document uploads, dashboard summaries, and report generation. This modular approach will make the system easier to test, maintain, and improve in future versions. Basic security practices such as input validation, password protection, role-based access, and controlled file uploads will also be applied to protect institutional records.")
    add_para(doc, "During development, each module will be checked against the identified requirements. The proponents will ensure that faculty users can submit and update their own records, while office personnel and administrators can validate entries and generate consolidated reports. The system layout will also be reviewed to make sure that users can navigate the platform easily and complete their tasks with minimal confusion.")

    add_section_heading(doc, "Testing")
    add_para(doc, "Testing will be conducted to determine whether the system performs its intended functions correctly and reliably. Unit testing will be used to check individual functions such as login validation, record saving, file upload handling, and report filtering. Integration testing will verify whether the modules work together properly, particularly the connection between proposal submission, document upload, status monitoring, and report generation.")
    add_para(doc, "System testing will evaluate the complete workflow of the application from the perspective of each user role. Faculty users will test the submission and updating of research and extension records, while Research and Extension Office personnel will test validation, monitoring, and report generation. Administrators will test account management, record access, and dashboard summaries. Security and usability checks will also be conducted to ensure that unauthorized users cannot access restricted data and that the interface remains understandable to intended users.")
    add_para(doc, "User Acceptance Testing will be conducted with selected representatives from the Research and Extension Office and faculty users. Their feedback will help determine whether the MIS is useful, easy to use, and aligned with the actual workflow of JHCSC Dumingag Campus. The results of testing will guide the correction of errors and the refinement of system features before full implementation.")

    add_section_heading(doc, "Implementation Plan")
    add_para(doc, "The implementation of the Web-Based MIS for Research and Extension will follow a phased approach. The first phase is preparation, which includes configuring the server environment, installing the required software, creating the database, and preparing the initial user accounts. During this stage, the proponents will also review the system settings and verify that the modules are ready for actual use.")
    add_para(doc, "The second phase is data preparation and migration. Available research and extension records will be reviewed, cleaned, and encoded into the system. Since the system cannot automatically digitize paper-based files, existing documents must be scanned or converted into digital format before uploading. This process will help establish an initial database that can support immediate retrieval and reporting.")
    add_para(doc, "The third phase is user training and orientation. Faculty members, office personnel, and administrators will be introduced to the major features of the system, including account login, record submission, document upload, status monitoring, budget entry, and report generation. Short guides may be provided to help users remember the steps needed for common tasks.")
    add_para(doc, "The final phase is pilot deployment and post-implementation support. The system will first be used by selected users to observe its performance in a real working environment. Any minor errors or usability concerns will be recorded and corrected. After the pilot period, the system may be adopted as the official platform for managing research and extension records at JHCSC Dumingag Campus.")

    add_chapter(doc, "CHAPTER V", "RESULTS, DISCUSSION, CONCLUSION AND RECOMMENDATION")
    add_section_heading(doc, "Results and Discussion")
    add_para(doc, "The Web-Based Management Information System for Research and Extension was developed to address the documentation and monitoring challenges encountered by J.H. Cerilles State College - Dumingag Campus. The system responds to the need for a centralized platform where research outputs, extension activities, supporting documents, budget records, and reports can be stored, updated, and retrieved efficiently. The results are discussed according to the objectives of the study.")
    add_para(doc, "First, the system supports the digital recording and monitoring of research outputs such as completed, ongoing, presented, published, and utilized research projects. Instead of relying on dispersed paper files or separate spreadsheets, faculty users can encode research details in a structured format and upload supporting documents. This improves the accuracy of records and allows office personnel to retrieve research information more quickly.")
    add_para(doc, "Second, the system provides an extension activity module that documents planned target activities, actual accomplishments, attendance, budget utilization, materials needed, and remarks for unaccomplished targets. This module helps the Research and Extension Office compare planned outputs with actual results and identify activities that require follow-up. It also improves transparency in monitoring extension project implementation.")
    add_para(doc, "Third, the MIS provides centralized document management for abstracts, certificates of presentation, conference proceedings, publication details, narrative reports, and other supporting files. By storing documents in one organized platform, the system reduces the risk of misplaced records and makes it easier for authorized users to prepare evidence for accreditation, institutional assessment, and reporting.")
    add_para(doc, "Fourth, system testing and validation confirm whether the modules function according to the needs of faculty members, researchers, administrators, and office personnel. The testing process checks the completeness of records, correctness of report outputs, accessibility of uploaded documents, and usability of the interface. Feedback from users helps refine the system and ensures that it supports actual work practices.")
    add_para(doc, "Lastly, the automated report generation feature improves administrative decision-making by producing summaries of research and extension records. Reports can support accreditation requirements, planning activities, and institutional monitoring. Overall, the system contributes to improved data accuracy, accessibility, and transparency in managing research and extension activities at JHCSC Dumingag Campus.")

    add_section_heading(doc, "Conclusions")
    add_para(doc, "Based on the development and evaluation of the Web-Based Management Information System for Research and Extension, the project addresses the major limitations of the existing manual and fragmented recordkeeping process. The system provides a centralized platform for recording, monitoring, storing, and retrieving research and extension data. It helps reduce paperwork, minimize duplicated records, and improve the organization of institutional documents.")
    add_para(doc, "The system also supports the needs of different users through role-based access. Faculty members can manage their own research and extension submissions, while the Research and Extension Office and administrators can validate records, monitor project status, and generate reports. These functions promote accountability, transparency, and more efficient management of research and extension activities.")
    add_para(doc, "In conclusion, the MIS strengthens the documentation process of JHCSC Dumingag Campus by making research and extension records more accessible, accurate, and useful for decision-making. With proper implementation and continued maintenance, the system can become a valuable tool for institutional assessment, accreditation preparation, and long-term monitoring of academic and community service outputs.")

    add_section_heading(doc, "Recommendations")
    add_para(doc, "Based on the results and conclusions of the study, the following recommendations are offered for future improvement of the Web-Based Management Information System for Research and Extension.")
    add_para(doc, "First, it is recommended to include an email or SMS notification feature in future versions of the system. This will help inform faculty members and administrators about proposal updates, validation results, report deadlines, and required corrections without requiring users to check the system manually at all times.")
    add_para(doc, "Second, the system may be expanded to support multi-campus integration if the institution decides to centralize research and extension monitoring across other campuses of J.H. Cerilles State College. This enhancement would allow higher-level administrators to view consolidated reports while still maintaining campus-specific records.")
    add_para(doc, "Third, the reporting module may be improved by adding graphical dashboards and data analytics. Charts showing research productivity, extension accomplishment rates, budget utilization, publication status, and project timelines would help administrators interpret data more quickly and make better planning decisions.")
    add_para(doc, "Fourth, regular backup procedures and stronger data security measures should be implemented. Since the system will store institutional documents and user information, scheduled backups, secure file handling, password policies, and access monitoring should be maintained to protect the integrity and confidentiality of records.")
    add_para(doc, "Lastly, future researchers may enhance the system by developing a mobile-friendly or progressive web application version. This would allow faculty members and office personnel to submit updates, upload documents, and check project status more conveniently using smartphones or tablets, especially during field-based extension activities.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
